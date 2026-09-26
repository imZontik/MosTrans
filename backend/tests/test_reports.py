from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.business.services.reports import _weakest_category
from app.frameworks.xlsx import training_report_xlsx


def run(category: str, outcome: str) -> dict:
    return {"category": category, "outcome": outcome}


def test_weakest_category_needs_enough_runs():
    # One failed medical run is not a gap yet; two conflicts with one failure are
    runs = [run("medical", "fail"), run("conflict", "fail"), run("conflict", "success")]
    assert _weakest_category(runs) == "Конфликты с пассажирами"


def test_no_weak_category_when_everything_passed():
    assert _weakest_category([run("safety", "success"), run("safety", "success")]) is None


def test_training_report_xlsx():
    now = datetime(2026, 9, 26, 12, 0, tzinfo=timezone.utc)
    employee = {
        "id": 1, "full_name": "Иван Петров", "email": "i@m400.ru", "position": "Проводник", "team": "Бригада 1",
        "level": 2, "points_total": 300, "points_period": 120, "runs": 4, "success_rate": 0.75, "avg_loyalty": 70.0,
        "avg_safety": 82.5, "safety_trend": -4.0, "timeout_rate": 0.1, "weak_category": None, "last_active_at": now,
    }
    report = {
        "generated_at": now,
        "days": 30,
        "since": now,
        "summary": {
            "employees": 1, "active": 1, "runs": 4, "success_rate": 0.75, "avg_loyalty": 70.0, "avg_safety": 82.5,
            "timeout_rate": 0.1, "emergencies": 0, "points": 120,
        },
        "employees": [employee],
        "categories": [],
        "mistakes": [],
    }
    wb = load_workbook(BytesIO(training_report_xlsx(report)))
    assert wb.sheetnames == ["Сводка", "Сотрудники", "Компетенции", "Частые ошибки"]
    people = wb["Сотрудники"]
    assert people["A1"].value == "Сотрудник"
    assert people["A2"].value == "Иван Петров"
    assert people["I2"].number_format == "0%"
    # Moscow time, no tzinfo: Excel has no time zones
    assert people["O2"].value == datetime(2026, 9, 26, 15, 0)
