"""Training report as an Excel workbook (openpyxl): what HR opens, filters and forwards."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.business.tournament import MSK

HEADER_FILL = PatternFill("solid", fgColor="E21A1A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
PERCENT = "0%"
SIGNED = "+0.0;-0.0;0"


def _local(moment: datetime | None) -> datetime | None:
    # Excel has no time zones: write Moscow time without tzinfo
    return moment.astimezone(MSK).replace(tzinfo=None) if moment else None


def _table(ws: Worksheet, columns: list[tuple[str, str, int, str | None]], rows: list[dict]) -> None:
    """columns: (key, title, width, number format)."""
    ws.append([title for _, title, _, _ in columns])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([row.get(key) for key, *_ in columns])
    for i, (_, _, width, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def training_report_xlsx(report: dict) -> bytes:
    wb = Workbook()
    s = report["summary"]

    ws = wb.active
    ws.title = "Сводка"
    ws.append(["Отчёт по обучению проводников ВСМ"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Период: {report['days']} дней, с {_local(report['since']):%d.%m.%Y}"])
    ws.append([f"Сформирован: {_local(report['generated_at']):%d.%m.%Y %H:%M} (МСК)"])
    ws.append([])
    for label, value, fmt in [
        ("Сотрудников", s["employees"], None),
        ("Тренировались за период", s["active"], None),
        ("Прохождений", s["runs"], None),
        ("Успешных прохождений", s["success_rate"], PERCENT),
        ("Средняя лояльность пассажира", s["avg_loyalty"], None),
        ("Средний рейтинг безопасности", s["avg_safety"], None),
        ("Доля решений, просроченных по таймеру", s["timeout_rate"], PERCENT),
        ("Специвентов пройдено", s["emergencies"], None),
        ("Начислено очков компетенций", s["points"], None),
    ]:
        ws.append([label, value])
        if fmt:
            ws.cell(ws.max_row, 2).number_format = fmt
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14

    employees = [{**e, "last_active_at": _local(e["last_active_at"])} for e in report["employees"]]
    _table(
        wb.create_sheet("Сотрудники"),
        [
            ("full_name", "Сотрудник", 30, None),
            ("position", "Должность", 20, None),
            ("team", "Бригада", 34, None),
            ("depot", "Депо", 22, None),
            ("level", "Уровень", 9, None),
            ("points_total", "Очки всего", 11, None),
            ("points_period", "Очки за период", 11, None),
            ("runs", "Прохождений", 12, None),
            ("success_rate", "Успешных", 10, PERCENT),
            ("avg_loyalty", "Лояльность", 11, None),
            ("avg_safety", "Безопасность", 13, None),
            ("safety_trend", "Динамика безопасности", 13, SIGNED),
            ("timeout_rate", "Таймауты", 10, PERCENT),
            ("weak_category", "Проседающая компетенция", 28, None),
            ("last_active_at", "Последняя активность", 18, "dd.mm.yyyy hh:mm"),
            ("email", "Почта", 24, None),
            ("id", "ID", 6, None),
        ],
        employees,
    )

    _table(
        wb.create_sheet("Компетенции"),
        [
            ("title", "Компетенция", 30, None),
            ("runs", "Прохождений", 12, None),
            ("employees", "Сотрудников", 12, None),
            ("success_rate", "Успешных", 10, PERCENT),
            ("avg_loyalty", "Лояльность", 11, None),
            ("avg_safety", "Безопасность", 13, None),
            ("timeout_rate", "Таймауты", 10, PERCENT),
        ],
        report["categories"],
    )

    mistakes = wb.create_sheet("Частые ошибки")
    _table(
        mistakes,
        [
            ("scenario", "Сценарий", 28, None),
            ("prompt", "Ситуация", 60, None),
            ("count", "Ошибок", 9, None),
            ("employees", "Сотрудников", 12, None),
            ("typical_answer", "Типичный неверный ответ", 50, None),
        ],
        report["mistakes"],
    )
    for row in mistakes.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
